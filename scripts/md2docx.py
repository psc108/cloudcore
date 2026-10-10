#!/usr/bin/env python3
"""Markdown (the subset our plans and reports use) to .docx in the house style:
Arial, navy/blue/green palette, section banners, colour-coded tables,
Courier New code, a callout box, page numbers.

Usage: python3 scripts/md2docx.py <in.md> <out.docx>   (needs python-docx)
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY, BLUE, TEXT = "1F3864", "2E4D7B", "1A1A1A"
LIGHT, GREY, GREEN, GREEN_LIGHT, AMBER_LIGHT = "D6E4F0", "F2F2F2", "375623", "E2EFDA", "FFF2CC"


def shade(cell_or_par, fill: str) -> None:
    el = cell_or_par._tc if hasattr(cell_or_par, "_tc") else cell_or_par._p
    pr = el.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), fill)
    pr.append(shd)


def runs(par, text: str, size: float = 10, color: str = TEXT, bold: bool = False) -> None:
    """Inline **bold**, *italic* and `code`."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
        if not part:
            continue
        b, i, code = bold, False, False
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            part, code = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        r = par.add_run(part.replace("<br>", "\n"))
        r.bold, r.italic = b, i
        r.font.name = "Courier New" if code else "Arial"
        r.font.size = Pt(size - 0.5 if code else size)
        r.font.color.rgb = RGBColor.from_string(color)


def banner(doc, text: str) -> None:
    p = doc.add_paragraph()
    shade(p, NAVY)
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(14), Pt(6)
    runs(p, text.upper(), 11, "FFFFFF", True)


def table(doc, rows: list[list[str]]) -> None:
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    result_cols = {n for n, h in enumerate(header) if h.strip("* ").lower() in ("result", "guard now", "harm")}
    for c, h in enumerate(header):
        cell = t.cell(0, c)
        shade(cell, NAVY)
        cell.paragraphs[0].text = ""
        runs(cell.paragraphs[0], h.replace("**", ""), 9, "FFFFFF", True)
    for r, row in enumerate(body, start=1):
        total = row[0].startswith("**Total")
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.paragraphs[0].text = ""
            runs(cell.paragraphs[0], val, 9, TEXT, total)
            if total:
                shade(cell, LIGHT)
            elif c in result_cols:
                shade(cell, AMBER_LIGHT if header[c].strip("* ").lower() == "harm" else GREEN_LIGHT)
            elif r % 2 == 0:
                shade(cell, GREY)
    if len(header) >= 6:
        # A wide numeric table: give the label column room, share out the rest.
        t.autofit = False
        rest = (16.6 - 4.2) / (len(header) - 1)
        for c, col in enumerate(t.columns):
            col.width = Cm(4.2 if c == 0 else rest)
        for row in t.rows:
            for c, cell in enumerate(row.cells):
                cell.width = Cm(4.2 if c == 0 else rest)
    doc.add_paragraph()


def callout(doc, text: str) -> None:
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    shade(cell, LIGHT)
    cell.paragraphs[0].text = ""
    runs(cell.paragraphs[0], text, 10, NAVY)
    doc.add_paragraph()


def page_numbers(doc) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for label, field in (("Page ", "PAGE"), (" of ", "NUMPAGES")):
        runs(p, label, 8, BLUE)
        r = p.add_run()
        for kind, txt in (("begin", None), (None, field), ("end", None)):
            if kind:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                r._r.append(fc)
            else:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = txt
                r._r.append(it)
        r.font.size = Pt(8)


def convert(src: Path, dst: Path, callout_prefix: str = "**The main lesson:**") -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(10)
    page_numbers(doc)
    lines = src.read_text().splitlines()
    i, section = 0, 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("# "):
            p = doc.add_paragraph()
            runs(p, line[2:], 26, NAVY, True)
            i += 1
            continue
        if line.startswith("## "):
            title = re.sub(r"^\d+\.\s*", "", line[3:])
            if title.lower() == "document history":
                banner(doc, "Document history")
            else:
                section += 1
                banner(doc, f"Section {section} — {title}")
            i += 1
            continue
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            shade(p, "F4F4F4")
            r = p.add_run("\n".join(code))
            r.font.name, r.font.size = "Courier New", Pt(9)
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table(doc, rows)
            continue
        m = re.match(r"^(\s*)- (.*)", line)
        if m:
            text = m.group(2)
            if text.startswith(callout_prefix):
                callout(doc, text)
            else:
                p = doc.add_paragraph(style="List Bullet 2" if len(m.group(1)) >= 2 else "List Bullet")
                runs(p, text)
            i += 1
            continue
        if line.strip():
            p = doc.add_paragraph()
            runs(p, line)
        i += 1
    doc.save(dst)


if __name__ == "__main__":
    convert(Path(sys.argv[1]), Path(sys.argv[2]))
